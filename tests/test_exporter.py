"""Tests for :mod:`sec_research_agent.reports.exporter`."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from sec_research_agent.reports.exporter import export_report, markdown_to_docx, markdown_to_pdf

SAMPLE_MARKDOWN = """\
# Example Report

## Overview

This is **bold** and this is *italic* text about the company.

- First bullet point
- Second bullet point

1. First step
2. Second step

---

A table:

| Metric | Value |
| --- | --- |
| Revenue | $1B |

```
code block line
```
"""


def test_export_report_writes_all_three_files(tmp_path: Path) -> None:
    files = export_report(
        SAMPLE_MARKDOWN, tmp_path, "AAPL_report_20240101_000000", brand_name="Involabs Financial Agent"
    )
    assert files.markdown.exists()
    assert files.docx.exists()
    assert files.pdf.exists()
    assert files.markdown.read_text(encoding="utf-8") == SAMPLE_MARKDOWN
    assert files.pdf.read_bytes().startswith(b"%PDF")
    assert files.docx.stat().st_size > 0


def test_markdown_to_docx_bold_and_headings(tmp_path: Path) -> None:
    path = markdown_to_docx(SAMPLE_MARKDOWN, tmp_path / "out.docx", brand_name="Involabs Financial Agent")
    doc = DocxDocument(str(path))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Example Report" in headings
    assert "Overview" in headings

    bold_paragraph = next(p for p in doc.paragraphs if "bold" in p.text)
    bold_runs = [r for r in bold_paragraph.runs if r.text == "bold"]
    assert bold_runs and bold_runs[0].bold is True

    italic_runs = [r for r in bold_paragraph.runs if r.text == "italic"]
    assert italic_runs and italic_runs[0].italic is True


def test_markdown_to_docx_lists(tmp_path: Path) -> None:
    path = markdown_to_docx(SAMPLE_MARKDOWN, tmp_path / "out.docx")
    doc = DocxDocument(str(path))
    bullet_paras = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    number_paras = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert len(bullet_paras) == 2
    assert len(number_paras) == 2


def test_markdown_to_docx_brand_header_optional(tmp_path: Path) -> None:
    with_brand = markdown_to_docx(
        "# T\n\ntext", tmp_path / "brand.docx", brand_name="Involabs Financial Agent"
    )
    without_brand = markdown_to_docx("# T\n\ntext", tmp_path / "no_brand.docx", brand_name="")
    doc_with = DocxDocument(str(with_brand))
    doc_without = DocxDocument(str(without_brand))
    assert doc_with.sections[0].header.paragraphs[0].runs
    assert not doc_without.sections[0].header.paragraphs[0].runs


def test_markdown_to_pdf_produces_valid_pdf_without_letterhead(tmp_path: Path) -> None:
    path = markdown_to_pdf(SAMPLE_MARKDOWN, tmp_path / "out.pdf", brand_name="Involabs Financial Agent")
    assert path.read_bytes().startswith(b"%PDF")


def test_markdown_to_pdf_missing_letterhead_image_falls_back_gracefully(tmp_path: Path) -> None:
    missing = tmp_path / "does_not_exist.png"
    path = markdown_to_pdf(
        SAMPLE_MARKDOWN, tmp_path / "out.pdf", brand_name="Involabs", letterhead_image=missing
    )
    assert path.read_bytes().startswith(b"%PDF")


def test_markdown_to_pdf_empty_document_does_not_crash(tmp_path: Path) -> None:
    path = markdown_to_pdf("", tmp_path / "empty.pdf")
    assert path.read_bytes().startswith(b"%PDF")


def test_export_report_creates_output_directory(tmp_path: Path) -> None:
    out_dir = tmp_path / "nested" / "reports"
    files = export_report("# Report\n\ntext", out_dir, "base")
    assert files.markdown.parent == out_dir
